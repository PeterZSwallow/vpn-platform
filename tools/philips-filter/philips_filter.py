#!/usr/bin/env python3
"""
Philips Filter — GUI-программа для отбора строк из большого XLSX-файла.

Ищет во всех листах строки, в любой ячейке которых встречается
Philips / Phillips / Филипс / Филлипс (без учёта регистра, в т.ч. внутри
слов: «PHILIPS», «Филипса», «Philips-123»), и сохраняет только эти строки
в Philips.xlsx.

Большие файлы читаются потоково (openpyxl read_only / write_only),
поэтому весь файл в память не загружается.

Запуск:  python philips_filter.py
Зависимости:  pip install openpyxl
"""

import os
import re
import threading
import queue

from openpyxl import Workbook, load_workbook

# Phil(l)ips / Фил(л)ипс, без учёта регистра
PATTERN = re.compile(r"phill?ips|филл?ипс", re.IGNORECASE)

OUTPUT_NAME = "Philips.xlsx"


def row_matches(row):
    for value in row:
        if value is not None and PATTERN.search(str(value)):
            return True
    return False


def filter_workbook(src_path, dst_path, keep_header=True, progress=None, stop_event=None):
    """
    Копирует из src_path в dst_path только строки с упоминанием Philips.
    Каждый лист исходника с найденными строками -> одноимённый лист в результате.
    progress(sheet_name, rows_scanned, rows_found) вызывается периодически.
    Возвращает (всего просмотрено строк, всего найдено строк).
    """
    src = load_workbook(src_path, read_only=True, data_only=True)
    dst = Workbook(write_only=True)

    total_scanned = 0
    total_found = 0

    try:
        for ws in src.worksheets:
            out_ws = None
            header = None
            scanned = found = 0

            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if stop_event is not None and stop_event.is_set():
                    raise InterruptedError("Остановлено пользователем")

                scanned += 1
                if i == 0 and keep_header:
                    header = row
                    if not row_matches(row):
                        continue

                if row_matches(row):
                    if out_ws is None:
                        out_ws = dst.create_sheet(title=ws.title)
                        if header is not None and i != 0:
                            out_ws.append(header)
                    out_ws.append(row)
                    found += 1

                if progress is not None and scanned % 5000 == 0:
                    progress(ws.title, total_scanned + scanned, total_found + found)

            total_scanned += scanned
            total_found += found
            if progress is not None:
                progress(ws.title, total_scanned, total_found)
    finally:
        src.close()

    if total_found == 0:
        # пустая книга с одним листом, чтобы файл был валидным
        empty = dst.create_sheet(title="Philips")
        empty.append(["Строки с Philips / Филипс не найдены"])

    dst.save(dst_path)
    return total_scanned, total_found


# ---------------------------------------------------------------- GUI

def run_gui():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox

    root = tk.Tk()
    root.title("Philips Filter — отбор строк из XLSX")
    root.geometry("640x260")
    root.minsize(520, 240)

    src_var = tk.StringVar()
    dst_var = tk.StringVar()
    header_var = tk.BooleanVar(value=True)
    status_var = tk.StringVar(value="Выберите входной XLSX-файл")

    events = queue.Queue()
    stop_event = threading.Event()

    def choose_src():
        path = filedialog.askopenfilename(
            title="Входной файл",
            filetypes=[("Excel", "*.xlsx *.xlsm"), ("Все файлы", "*.*")],
        )
        if path:
            src_var.set(path)
            dst_var.set(os.path.join(os.path.dirname(path), OUTPUT_NAME))

    def choose_dst():
        path = filedialog.asksaveasfilename(
            title="Сохранить как",
            defaultextension=".xlsx",
            initialfile=OUTPUT_NAME,
            initialdir=os.path.dirname(dst_var.get() or src_var.get() or "."),
            filetypes=[("Excel", "*.xlsx")],
        )
        if path:
            dst_var.set(path)

    def worker(src, dst, keep_header):
        def progress(sheet, scanned, found):
            events.put(("progress", sheet, scanned, found))

        try:
            scanned, found = filter_workbook(src, dst, keep_header, progress, stop_event)
            events.put(("done", scanned, found, dst))
        except InterruptedError as e:
            events.put(("error", str(e)))
        except Exception as e:  # noqa: BLE001 — показываем любую ошибку пользователю
            events.put(("error", f"{type(e).__name__}: {e}"))

    def start():
        src, dst = src_var.get().strip(), dst_var.get().strip()
        if not src or not os.path.isfile(src):
            messagebox.showwarning("Нет файла", "Выберите существующий входной XLSX-файл.")
            return
        if not dst:
            dst = os.path.join(os.path.dirname(src), OUTPUT_NAME)
            dst_var.set(dst)
        if os.path.abspath(src) == os.path.abspath(dst):
            messagebox.showwarning("Ошибка", "Файл вывода совпадает с входным файлом.")
            return

        stop_event.clear()
        start_btn.config(state="disabled")
        stop_btn.config(state="normal")
        bar.start(12)
        status_var.set("Чтение файла…")
        threading.Thread(target=worker, args=(src, dst, header_var.get()), daemon=True).start()

    def stop():
        stop_event.set()
        status_var.set("Остановка…")

    def finish():
        bar.stop()
        start_btn.config(state="normal")
        stop_btn.config(state="disabled")

    def poll():
        try:
            while True:
                ev = events.get_nowait()
                if ev[0] == "progress":
                    _, sheet, scanned, found = ev
                    status_var.set(f"Лист «{sheet}»: просмотрено {scanned:,} строк, найдено {found:,}"
                                   .replace(",", " "))
                elif ev[0] == "done":
                    _, scanned, found, dst = ev
                    finish()
                    msg = f"Готово. Просмотрено {scanned:,} строк, найдено {found:,}.".replace(",", " ")
                    status_var.set(msg)
                    messagebox.showinfo("Готово", f"{msg}\n\nСохранено в:\n{dst}")
                elif ev[0] == "error":
                    finish()
                    status_var.set(ev[1])
                    messagebox.showerror("Ошибка", ev[1])
        except queue.Empty:
            pass
        root.after(150, poll)

    frm = ttk.Frame(root, padding=12)
    frm.pack(fill="both", expand=True)
    frm.columnconfigure(1, weight=1)

    ttk.Label(frm, text="Входной файл:").grid(row=0, column=0, sticky="w", pady=4)
    ttk.Entry(frm, textvariable=src_var).grid(row=0, column=1, sticky="ew", padx=6)
    ttk.Button(frm, text="Обзор…", command=choose_src).grid(row=0, column=2)

    ttk.Label(frm, text="Файл вывода:").grid(row=1, column=0, sticky="w", pady=4)
    ttk.Entry(frm, textvariable=dst_var).grid(row=1, column=1, sticky="ew", padx=6)
    ttk.Button(frm, text="Обзор…", command=choose_dst).grid(row=1, column=2)

    ttk.Checkbutton(frm, text="Сохранять первую строку (заголовок) каждого листа",
                    variable=header_var).grid(row=2, column=0, columnspan=3, sticky="w", pady=6)

    ttk.Label(frm, text="Ищется: Philips, Phillips, Филипс, Филлипс (без учёта регистра)",
              foreground="gray").grid(row=3, column=0, columnspan=3, sticky="w")

    bar = ttk.Progressbar(frm, mode="indeterminate")
    bar.grid(row=4, column=0, columnspan=3, sticky="ew", pady=10)

    ttk.Label(frm, textvariable=status_var).grid(row=5, column=0, columnspan=3, sticky="w")

    btns = ttk.Frame(frm)
    btns.grid(row=6, column=0, columnspan=3, sticky="e", pady=(10, 0))
    start_btn = ttk.Button(btns, text="Найти и сохранить", command=start)
    start_btn.pack(side="left", padx=4)
    stop_btn = ttk.Button(btns, text="Стоп", command=stop, state="disabled")
    stop_btn.pack(side="left")

    root.after(150, poll)
    root.mainloop()


if __name__ == "__main__":
    run_gui()
