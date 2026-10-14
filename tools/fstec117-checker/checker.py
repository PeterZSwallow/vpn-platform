"""Проверка технического паспорта объекта информатизации (DOCX)
на соответствие Приказу ФСТЭК России от 11.04.2025 № 117.

Логика не зависит от GUI и может использоваться из командной строки:
    python checker.py паспорт.docx [результат.docx]
"""

from __future__ import annotations

import copy
import datetime as dt
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from rules import OBSOLETE, P117, P117_FULL, PLACEHOLDERS, RULES

RED = RGBColor(0xFF, 0x00, 0x00)
FLAGS = re.IGNORECASE | re.UNICODE
NEIGHBOURS = 3  # сколько следующих абзацев смотреть в поисках значения

KIND_TITLES = {
    "missing": "Отсутствует раздел",
    "no_value": "Не указано значение",
    "obsolete": "Устаревшая ссылка",
    "placeholder": "Не заполнено",
    "empty_cell": "Пустая ячейка",
    "expired": "Истёк срок",
}


@dataclass
class Finding:
    kind: str
    level: str
    title: str
    message: str
    ref: str = ""
    context: str = ""
    anchor: Paragraph | None = field(default=None, repr=False)
    cell: object | None = field(default=None, repr=False)

    @property
    def note(self) -> str:
        text = f"[ПРОВЕРКА 117] {self.message}"
        if self.ref:
            text += f" — ссылка: {self.ref}"
        return text


@dataclass
class Block:
    paragraph: Paragraph
    row_text: str  # текст всей строки таблицы (или самого абзаца)
    in_table: bool


# ---------------------------------------------------------------- обход документа

def _iter_tables(tables, seen):
    for table in tables:
        for row in table.rows:
            row_cells = []
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                row_cells.append(cell)
            yield row, row_cells
            for cell in row_cells:
                yield from _iter_tables(cell.tables, seen)


def iter_blocks(doc) -> list[Block]:
    """Все абзацы документа в порядке следования, включая ячейки таблиц."""
    blocks: list[Block] = []
    body = doc.element.body
    seen: set[int] = set()
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, doc)
            blocks.append(Block(p, p.text, False))
        elif child.tag == qn("w:tbl"):
            for _row, cells in _iter_tables([Table(child, doc)], seen):
                row_text = " | ".join(c.text for c in cells)
                for cell in cells:
                    for p in cell.paragraphs:
                        blocks.append(Block(p, row_text, True))
    return blocks


def _search(patterns, text):
    return any(re.search(p, text, FLAGS) for p in patterns)


def _short(text, n=120):
    text = " ".join(text.split())
    return text if len(text) <= n else text[: n - 1] + "…"


# ---------------------------------------------------------------- проверки

DATE_RE = re.compile(
    r"(?:до|по|истекает|действ\w+\s+до)\s*(\d{2})\.(\d{2})\.(\d{4})", FLAGS)


def check_document(doc, today: dt.date | None = None) -> list[Finding]:
    today = today or dt.date.today()
    blocks = iter_blocks(doc)
    findings: list[Finding] = []

    # 1. Обязательные разделы и их значения.
    for rule in RULES:
        idx = next((i for i, b in enumerate(blocks)
                    if _search(rule["patterns"], b.paragraph.text)), None)
        if idx is None:
            findings.append(Finding(
                "missing", rule["level"], rule["title"],
                f"ТУТ ДОЛЖНО БЫТЬ: {rule['what']}", rule["ref"]))
            continue
        values = rule.get("value_patterns")
        if not values:
            continue
        # Значение ищем во всех местах, где упоминается раздел.
        ok = False
        first = blocks[idx]
        for i, b in enumerate(blocks):
            if not _search(rule["patterns"], b.paragraph.text):
                continue
            around = [b.row_text] + [x.paragraph.text for x in blocks[i + 1:i + 1 + NEIGHBOURS]]
            if _search(values, " \n".join(around)):
                ok = True
                break
        if not ok:
            findings.append(Finding(
                "no_value", rule["level"], rule["title"],
                f"ТУТ ДОЛЖНО БЫТЬ: {rule['what']}", rule["ref"],
                _short(first.row_text), first.paragraph))

    # 2. Абзацные проверки: устаревшие ссылки, заглушки, сроки.
    for b in blocks:
        text = b.paragraph.text
        if not text.strip():
            continue
        for pattern, msg in OBSOLETE:
            if re.search(pattern, text, FLAGS):
                findings.append(Finding("obsolete", "critical", "Нормативная ссылка",
                                        msg, "", _short(text), b.paragraph))
        flagged = any(f.anchor is b.paragraph for f in findings)
        signature = re.search(r"утвержда|подпис|М\.\s*П\.|/\s*\S+.*/", text, FLAGS)
        if _search(PLACEHOLDERS, text) and not flagged and not signature:
            findings.append(Finding(
                "placeholder", "critical", "Незаполненное поле",
                "Поле не заполнено. ТУТ ДОЛЖНО БЫТЬ конкретное значение", P117,
                _short(text), b.paragraph))
        if re.search(r"сертификат|аттестат", b.row_text, FLAGS):
            for d, m, y in DATE_RE.findall(text):
                try:
                    date = dt.date(int(y), int(m), int(d))
                except ValueError:
                    continue
                if date < today:
                    findings.append(Finding(
                        "expired", "critical", "Срок действия",
                        f"Срок действия ({d}.{m}.{y}) истёк. ДОЛЖЕН БЫТЬ действующий "
                        "сертификат/аттестат или сведения о его продлении",
                        P117, _short(b.row_text), b.paragraph))

    # 3. Пустые ячейки «значений» в таблицах вида «параметр | значение».
    seen: set[int] = set()
    for row, cells in _iter_tables(doc.tables, seen):
        if len(cells) < 2:
            continue
        label = cells[0].text.strip()
        if not label:
            continue
        empty = [cell for cell in cells[1:] if not cell.text.strip()]
        if empty:
            findings.append(Finding(
                "empty_cell", "warning", _short(label, 60),
                f"ТУТ ДОЛЖНО БЫТЬ: значение параметра «{_short(label, 80)}»", "",
                _short(label), None, empty[0]))
    return findings


# ---------------------------------------------------------------- запись правок

def _red_run(paragraph: Paragraph, text: str, bold=True, size=None):
    run = paragraph.add_run(text)
    run.font.color.rgb = RED
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return run


def _insert_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = copy.deepcopy(paragraph._p)
    for child in list(new_p):
        if child.tag != qn("w:pPr"):
            new_p.remove(child)
    # нумерацию списка не наследуем
    ppr = new_p.find(qn("w:pPr"))
    if ppr is not None:
        num = ppr.find(qn("w:numPr"))
        if num is not None:
            ppr.remove(num)
    paragraph._p.addnext(new_p)
    para = Paragraph(new_p, paragraph._parent)
    _red_run(para, text)
    return para


def annotate(doc, findings: list[Finding], source_name: str = "") -> None:
    """Вносит правки красным в документ (изменяет doc на месте)."""
    # Инлайн-замечания рядом с местом ошибки.
    last_inserted: dict[int, Paragraph] = {}
    for f in findings:
        if f.cell is not None:
            p = f.cell.paragraphs[0]
            _red_run(p, f.note)
        elif f.anchor is not None:
            if f.kind == "obsolete":
                for run in f.anchor.runs:
                    run.font.color.rgb = RED
                    run.font.strike = True
            key = id(f.anchor._p)
            after = last_inserted.get(key, f.anchor)
            last_inserted[key] = _insert_after(after, f.note)

    # Итоговая сводка в конце документа.
    doc.add_page_break()
    head = doc.add_paragraph()
    _red_run(head, "ЗАМЕЧАНИЯ ПО РЕЗУЛЬТАТАМ ПРОВЕРКИ НА СООТВЕТСТВИЕ ПРИКАЗУ "
                   "ФСТЭК РОССИИ ОТ 11.04.2025 № 117", size=14)
    _red_run(doc.add_paragraph(), P117_FULL, bold=False, size=10)
    info = doc.add_paragraph()
    crit = sum(f.level == "critical" for f in findings)
    _red_run(info, f"Документ: {source_name or '—'}. Дата проверки: "
                   f"{dt.date.today():%d.%m.%Y}. Всего замечаний: {len(findings)} "
                   f"(обязательных: {crit}, рекомендаций: {len(findings) - crit}).", bold=False)
    if not findings:
        _red_run(doc.add_paragraph(), "Замечаний не выявлено.", bold=False)
    for n, f in enumerate(findings, 1):
        p = doc.add_paragraph()
        level = "ОБЯЗАТЕЛЬНО" if f.level == "critical" else "РЕКОМЕНДУЕТСЯ"
        _red_run(p, f"{n}. [{level}] {f.title} ({KIND_TITLES[f.kind]}): ")
        _red_run(p, f.message + (f" — ссылка: {f.ref}" if f.ref else ""), bold=False)
        if f.context:
            _red_run(p, f" (фрагмент: «{f.context}»)", bold=False)
    _red_run(doc.add_paragraph(),
             "Автоматическая проверка носит вспомогательный характер: выводы и номера "
             "пунктов Приказа № 117 подлежат сверке с официальным текстом специалистом "
             "по защите информации.", bold=False, size=9)

    # Пометка в начале документа.
    body = doc.element.body
    first = body[0] if len(body) else None
    if first is not None and first.tag == qn("w:p"):
        top = Paragraph(first, doc).insert_paragraph_before()
    else:
        top = doc.add_paragraph()
        body.remove(top._p)
        body.insert(0, top._p)
    _red_run(top, f"[ПРОВЕРКА 117] Выявлено замечаний: {len(findings)}. Правки выделены "
                  "красным, сводка — в конце документа.")


def process(src: str | Path, dst: str | Path | None = None):
    src = Path(src)
    doc = Document(str(src))
    findings = check_document(doc)
    if dst:
        annotate(doc, findings, src.name)
        doc.save(str(dst))
    return findings


def default_output(src: str | Path) -> Path:
    src = Path(src)
    return src.with_name(f"{src.stem}_проверка_117{src.suffix}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Использование: python checker.py паспорт.docx [результат.docx]")
        sys.exit(1)
    out = sys.argv[2] if len(sys.argv) > 2 else default_output(sys.argv[1])
    result = process(sys.argv[1], out)
    for i, f in enumerate(result, 1):
        print(f"{i:3}. [{f.level}] {f.title}: {f.message}")
    print(f"\nСохранено: {out}")
